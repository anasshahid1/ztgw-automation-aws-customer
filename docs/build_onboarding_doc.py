#!/usr/bin/env python3
"""Generate a professional customer onboarding guide for ZTGW AWS Terraform deployment."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# -----------------------------------------------------------------------
# Brand palette
# -----------------------------------------------------------------------
ORANGE = RGBColor(0xE8, 0x5C, 0x2E)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
CODE_BG = "2D2D2D"
CODE_FG = RGBColor(0xE8, 0xE8, 0xE8)
TABLE_HEADER_BG = "E85C2E"
CALLOUT_BG = "FFF4EF"

doc = Document()

# -----------------------------------------------------------------------
# Global style spacing -- this is the actual fix. Word's default template
# gives Normal/Heading/List styles generous space_before/space_after that
# compounds badly once we also add our own paragraphs around tables. We
# override every style used in this document to tight, explicit values.
# -----------------------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

heading_sizes = {1: 22, 2: 15, 3: 12.5}
heading_space_before = {1: 18, 2: 14, 3: 10}
heading_space_after = {1: 6, 2: 4, 3: 3}
for i in (1, 2, 3):
    h = doc.styles[f"Heading {i}"]
    h.font.name = "Calibri"
    h.font.size = Pt(heading_sizes[i])
    h.font.color.rgb = ORANGE
    h.font.bold = True
    h.paragraph_format.space_before = Pt(heading_space_before[i])
    h.paragraph_format.space_after = Pt(heading_space_after[i])
    h.paragraph_format.keep_with_next = True

for style_name in ("List Bullet", "List Number"):
    s = doc.styles[style_name]
    s.font.name = "Calibri"
    s.font.size = Pt(10.5)
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after = Pt(2)
    s.paragraph_format.line_spacing = 1.1


# -----------------------------------------------------------------------
# Helpers
# -----------------------------------------------------------------------
def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    """Margins in twentieths of a point (dxa). Word default is ~108 (~0.075in)."""
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def tighten(paragraph, before=0, after=0, line=1.0):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_code_block(document, code_text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    cell = table.rows[0].cells[0]
    set_cell_background(cell, CODE_BG)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    cell.paragraphs[0].text = ""
    lines = code_text.strip("\n").split("\n")
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = CODE_FG
        tighten(p, before=0, after=0, line=1.0)
    spacer = document.add_paragraph()
    tighten(spacer, before=0, after=8)


def add_callout(document, label, text, color=ORANGE):
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    set_cell_background(cell, CALLOUT_BG)
    set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
    p = cell.paragraphs[0]
    tighten(p, before=0, after=0, line=1.15)
    run = p.add_run(f"{label}  ")
    run.bold = True
    run.font.color.rgb = color
    run.font.size = Pt(10)
    run2 = p.add_run(text)
    run2.font.size = Pt(10)
    run2.font.color.rgb = DARK
    gap = document.add_paragraph()
    tighten(gap, before=0, after=8)


def add_bullet(document, text, bold_prefix=None):
    p = document.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def add_numbered(document, text):
    p = document.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_body(document, text):
    p = document.add_paragraph(text)
    tighten(p, before=0, after=8, line=1.15)
    return p


def add_spacer(document, pt=6):
    p = document.add_paragraph()
    tighten(p, before=0, after=pt)
    return p


def styled_table(document, headers, rows, col_widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        set_cell_margins(hdr_cells[i], top=50, bottom=50, left=90, right=90)
        p = hdr_cells[i].paragraphs[0]
        tighten(p, before=0, after=0, line=1.05)
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)
        set_cell_background(hdr_cells[i], TABLE_HEADER_BG)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            set_cell_margins(cells[i], top=50, bottom=50, left=90, right=90)
            p = cells[i].paragraphs[0]
            tighten(p, before=0, after=0, line=1.05)
            run = p.add_run(str(val))
            run.font.size = Pt(9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    gap = document.add_paragraph()
    tighten(gap, before=0, after=10)
    return table


def add_heading_tight(document, text, level):
    h = document.add_heading(text, level=level)
    return h


# -----------------------------------------------------------------------
# COVER PAGE
# -----------------------------------------------------------------------
for _ in range(5):
    add_spacer(doc, 14)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(title, before=0, after=4)
run = title.add_run("AWS Zero Trust Gateway (ZTGW)")
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = ORANGE

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(subtitle, before=0, after=10)
run = subtitle.add_run("Terraform Deployment — Customer Onboarding Guide")
run.font.size = Pt(17)
run.font.color.rgb = DARK

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(sub2, before=0, after=0, line=1.3)
run = sub2.add_run("Provisioning a Zscaler AWS Zero Trust Gateway and Gateway Load "
                   "Balancer Endpoint using Infrastructure as Code")
run.font.size = Pt(11.5)
run.font.italic = True
run.font.color.rgb = GRAY

for _ in range(7):
    add_spacer(doc, 14)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(meta, before=0, after=2)
run = meta.add_run("Document Version 1.0")
run.font.size = Pt(10.5)
run.font.color.rgb = GRAY

meta2 = doc.add_paragraph()
meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(meta2, before=0, after=0)
run = meta2.add_run("Prepared for Customer Deployment Teams")
run.font.size = Pt(10.5)
run.font.color.rgb = GRAY

doc.add_page_break()

# -----------------------------------------------------------------------
# TABLE OF CONTENTS
# -----------------------------------------------------------------------
add_heading_tight(doc, "Table of Contents", 1)
toc_items = [
    "1. Overview",
    "2. Architecture",
    "3. Prerequisites",
    "4. Deployment Method A — Pure Terraform Workflow",
    "5. Deployment Method B — Interactive zsec Wrapper",
    "6. Variable Reference",
    "7. Verifying Your Deployment",
    "8. Tearing Down / Destroying Resources",
    "9. Troubleshooting",
    "10. Security Notes",
]
for item in toc_items:
    p = doc.add_paragraph(item)
    tighten(p, before=0, after=5)

doc.add_page_break()

# -----------------------------------------------------------------------
# 1. OVERVIEW
# -----------------------------------------------------------------------
add_heading_tight(doc, "1. Overview", 1)
add_body(doc,
    "This guide walks through deploying an AWS Zero Trust Gateway (ZTGW) and its "
    "associated AWS Gateway Load Balancer (GWLB) Endpoint using Terraform. The ZTGW "
    "is provisioned via the Zscaler REST API, and the GWLB Endpoint is created "
    "directly as native Terraform resources inside your existing AWS VPC."
)
add_body(doc, "Two supported deployment methods are provided, and you may choose whichever best fits your team's workflow:")
add_bullet(doc, " A single terraform apply command drives the entire deployment "
                "from a pre-filled configuration file. Recommended for teams already "
                "comfortable with Terraform and CI/CD automation.",
           bold_prefix="Method A — Pure Terraform:")
add_bullet(doc, " A guided, interactive command-line wizard prompts you for each "
                "required value and caches your configuration for repeat runs. "
                "Recommended for first-time or manual deployments.",
           bold_prefix="Method B — zsec Interactive Wrapper:")
add_spacer(doc, 4)
add_callout(doc, "Note:", "Both methods produce the same end result and can be used "
                          "interchangeably across different environments.")

# -----------------------------------------------------------------------
# 2. ARCHITECTURE
# -----------------------------------------------------------------------
add_heading_tight(doc, "2. Architecture", 1)
add_body(doc, "At a high level, the deployment flow is as follows:")
flow_steps = [
    "Your configuration values (Zscaler credentials, AWS region, VPC/subnet IDs) are supplied via a configuration file.",
    "Terraform authenticates to the Zscaler OneAPI and creates the AWS Zero Trust Gateway.",
    "Terraform polls the gateway's health status until it reports Healthy (this can take several minutes).",
    "The gateway's VPC Endpoint Service Name is retrieved automatically.",
    "Terraform creates a Gateway Load Balancer (GWLB) Endpoint in each subnet/Availability Zone you specify, connected to that service.",
    "Terraform creates a route in each subnet's route table, directing traffic to its corresponding GWLB Endpoint for inspection.",
]
for s in flow_steps:
    add_numbered(doc, s)
add_spacer(doc, 4)
add_callout(doc, "Important:", "AWS requires exactly one subnet per Gateway Load "
                                "Balancer-type VPC Endpoint. If you provide subnets "
                                "across two Availability Zones, two separate endpoints "
                                "and two separate routes will be created — one per "
                                "subnet — to support multi-AZ redundancy.")

add_heading_tight(doc, "Deployment Modes", 2)
styled_table(
    doc,
    ["Mode", "Create Gateway?", "Create Endpoint?", "Use Case"],
    [
        ["Full", "Yes", "Yes", "Deploy everything in a single run (recommended for most customers)"],
        ["Gateway-only", "Yes", "No", "Create the gateway first; add endpoints in a later, separate run"],
        ["Endpoint-only", "No", "Yes", "Attach a new VPC/subnet to a gateway that already exists"],
    ],
    col_widths=[1.1, 1.1, 1.1, 2.9],
)


# -----------------------------------------------------------------------
# 3. PREREQUISITES
# -----------------------------------------------------------------------
add_heading_tight(doc, "3. Prerequisites", 1)
add_body(doc, "Before beginning, ensure the following are available:")

styled_table(
    doc,
    ["Requirement", "Purpose"],
    [
        ["Terraform >= 1.0", "Infrastructure orchestration"],
        ["AWS CLI", "Availability Zone discovery; AWS credential configuration"],
        ["Python 3", "Used internally by deployment scripts for JSON processing"],
        ["Zscaler OneAPI Credentials", "Client ID, Client Secret, and Vanity Domain for your Zscaler tenant"],
        ["AWS Credentials", "Configured for the target AWS account (via profile, environment variables, or IAM role)"],
        ["Existing VPC and Subnets", "At least one subnet per Availability Zone where the GWLB Endpoint(s) will be created"],
    ],
    col_widths=[2.2, 4.0],
)

add_heading_tight(doc, "Required AWS IAM Permissions", 2)
add_body(doc, "The AWS credentials used must be able to perform the following actions:")
for perm in [
    "ec2:CreateVpcEndpoint", "ec2:DeleteVpcEndpoints", "ec2:DescribeVpcEndpoints",
    "ec2:DescribeSubnets", "ec2:DescribeRouteTables", "ec2:CreateRoute",
    "ec2:DeleteRoute", "ec2:DescribeAvailabilityZones",
]:
    add_bullet(doc, perm)


# -----------------------------------------------------------------------
# 4. METHOD A — PURE TERRAFORM
# -----------------------------------------------------------------------
add_heading_tight(doc, "4. Deployment Method A — Pure Terraform Workflow", 1)
add_body(doc,
    "This method drives the entire deployment through a single Terraform "
    "configuration file and a single terraform apply command."
)

add_heading_tight(doc, "Step 1 — Clone the repository", 2)
add_code_block(doc, "git clone <repository-url>\ncd <repository-folder>")

add_heading_tight(doc, "Step 2 — Create your configuration file", 2)
add_body(doc,
    "Copy the example configuration file and fill in your real values. "
    "Never commit the resulting file to version control — it will contain "
    "sensitive credentials."
)
add_code_block(doc, "cp ztgw-full-deploy.tfvars.example ztgw-full-deploy.tfvars")

add_body(doc, "Edit ztgw-full-deploy.tfvars and supply the following:")
for item in [
    "Zscaler OneAPI client_id, client_secret, and vanity_domain",
    "Target AWS region and region code",
    "A unique gateway_name and location_name",
    "Your AWS account ID (to whitelist it for the endpoint connection)",
    "Your existing VPC ID and subnet IDs (one per Availability Zone)",
]:
    add_bullet(doc, item)

add_heading_tight(doc, "Step 3 — Initialize Terraform", 2)
add_code_block(doc, "terraform init")

add_heading_tight(doc, "Step 4 — Review the plan (optional but recommended)", 2)
add_code_block(doc, "terraform plan -var-file=ztgw-full-deploy.tfvars")

add_heading_tight(doc, "Step 5 — Apply the deployment", 2)
add_code_block(doc, "terraform apply -var-file=ztgw-full-deploy.tfvars")
add_body(doc,
    "Terraform will prompt for confirmation before proceeding. Type yes to continue. "
    "The deployment will then:"
)
for s in [
    "Create the Zero Trust Gateway via the Zscaler API",
    "Wait for the gateway to report a Healthy status (typically 3–10 minutes)",
    "Create the GWLB Endpoint(s) in your specified subnet(s)",
    "Create the corresponding route(s) in each subnet's route table",
]:
    add_bullet(doc, s)
add_spacer(doc, 4)
add_callout(doc, "Please be patient:", "Gateway provisioning can take several minutes. "
                                       "Progress is printed to your terminal throughout.")

add_heading_tight(doc, "Step 6 — Review the outputs", 2)
add_body(doc, "Upon successful completion, Terraform will display outputs including the gateway ID, "
              "health status, and VPC Endpoint details. You may retrieve these at any time with:")
add_code_block(doc, "terraform output")


# -----------------------------------------------------------------------
# 5. METHOD B — ZSEC
# -----------------------------------------------------------------------
add_heading_tight(doc, "5. Deployment Method B — Interactive zsec Wrapper", 1)
add_body(doc,
    "This method uses a guided command-line wizard that prompts for each "
    "required value, making it well suited for first-time or manual deployments."
)

add_heading_tight(doc, "Step 1 — Clone the repository", 2)
add_code_block(doc, "git clone <repository-url>\ncd <repository-folder>")

add_heading_tight(doc, "Step 2 — Run the interactive wizard", 2)
add_code_block(doc, "./zsec up")

add_body(doc, "On first run, you will be prompted for:")
for s in [
    "Deployment mode (Full / Gateway-only / Endpoint-only)",
    "Zscaler OneAPI credentials",
    "AWS region and gateway configuration",
    "Consumer VPC and subnet details (if deploying the endpoint)",
    "A final confirmation summary before anything is created",
]:
    add_numbered(doc, s)

add_body(doc,
    "Your responses are cached in a local configuration file so that re-running "
    "./zsec up in the future will skip the prompts and redeploy using the same values."
)

add_heading_tight(doc, "Step 3 — Monitor progress", 2)
add_body(doc,
    "The wizard will create the gateway, wait for it to become healthy, and then "
    "run Terraform automatically to create the VPC Endpoint(s) and routes — all "
    "within the same command."
)


# -----------------------------------------------------------------------
# 6. VARIABLE REFERENCE
# -----------------------------------------------------------------------
add_heading_tight(doc, "6. Variable Reference", 1)

add_heading_tight(doc, "Authentication", 2)
styled_table(
    doc,
    ["Variable", "Required", "Description"],
    [
        ["client_id", "Yes", "Zscaler OneAPI OAuth2 client ID"],
        ["client_secret", "Yes", "Zscaler OneAPI OAuth2 client secret"],
        ["vanity_domain", "Yes", "Your Zscaler tenant's vanity domain"],
        ["cloud", "Yes", "Zscaler cloud identifier (e.g. zscalerthree)"],
        ["login_domain", "No", "Override OAuth2 login domain if your tenant requires it"],
    ],
    col_widths=[1.8, 1.0, 3.4],
)

add_heading_tight(doc, "Gateway Configuration", 2)
styled_table(
    doc,
    ["Variable", "Required", "Description"],
    [
        ["gateway_name", "Yes", "Unique name for the Zero Trust Gateway"],
        ["location_name", "Yes", "Unique Zscaler location name associated with the gateway"],
        ["aws_region", "Yes", "AWS region for the gateway (e.g. us-west-2)"],
        ["aws_region_code", "Yes", "Zscaler-format region code (e.g. US_WEST_2)"],
        ["availability_zone_ids", "No", "AZ IDs; auto-discovered if left empty"],
        ["additional_aws_accounts", "No", "AWS account IDs permitted to connect to the endpoint service"],
    ],
    col_widths=[1.8, 1.0, 3.4],
)

add_heading_tight(doc, "Consumer VPC Endpoint", 2)
styled_table(
    doc,
    ["Variable", "Required", "Description"],
    [
        ["consumer_vpc_id", "Yes", "Existing VPC ID where the endpoint will be created"],
        ["consumer_subnet_ids", "Yes", "Existing subnet IDs, one per Availability Zone"],
        ["consumer_route_cidr", "No", "Destination CIDR for the inspection route (default: 0.0.0.0/0)"],
        ["consumer_deploy_key", "No", "Naming prefix applied to created resources"],
    ],
    col_widths=[1.8, 1.0, 3.4],
)


# -----------------------------------------------------------------------
# 7. VERIFYING YOUR DEPLOYMENT
# -----------------------------------------------------------------------
add_heading_tight(doc, "7. Verifying Your Deployment", 1)
add_body(doc, "After deployment completes, confirm success using the following checks:")

add_bullet(doc, " Confirm the gateway shows a Healthy status in Terraform outputs or the Zscaler Admin Portal.",
           bold_prefix="Gateway Health: ")
add_bullet(doc, " Confirm the VPC Endpoint(s) show an Available state in the AWS Console (VPC > Endpoints).",
           bold_prefix="VPC Endpoint State: ")
add_bullet(doc, " Confirm each subnet's route table contains a route to its corresponding VPC Endpoint.",
           bold_prefix="Route Tables: ")
add_bullet(doc, " Deploy a test workload in the protected subnet and confirm its traffic is visible for inspection in the Zscaler Admin Portal.",
           bold_prefix="Traffic Flow: ")


# -----------------------------------------------------------------------
# 8. DESTROY
# -----------------------------------------------------------------------
add_heading_tight(doc, "8. Tearing Down / Destroying Resources", 1)

add_heading_tight(doc, "Method A — Pure Terraform", 2)
add_code_block(doc, "terraform destroy -var-file=ztgw-full-deploy.tfvars")
add_body(doc,
    "This removes all resources in the correct order: routes, then VPC Endpoints, "
    "then the Zero Trust Gateway itself."
)

add_heading_tight(doc, "Method B — zsec", 2)
add_code_block(doc, "./zsec destroy")
add_body(doc,
    "This performs the same cleanup via the interactive wrapper, with a "
    "confirmation prompt before proceeding."
)

add_callout(doc, "Important:", "Deleting a gateway does not automatically remove its "
                                "associated Zscaler Location object. If you plan to "
                                "redeploy using the same location name, see the "
                                "Troubleshooting section below.")


# -----------------------------------------------------------------------
# 9. TROUBLESHOOTING
# -----------------------------------------------------------------------
add_heading_tight(doc, "9. Troubleshooting", 1)

issues = [
    (
        "Gateway creation fails with a generic API validation error",
        "Double-check every configuration value against the Variable Reference "
        "section. A common cause is placing an AWS account ID into the wrong "
        "field — AWS account IDs belong in additional_aws_accounts, not in "
        "Zscaler-internal account/group fields.",
    ),
    (
        "Location name already exists",
        "Each gateway creates an associated Zscaler Location object. Deleting a "
        "gateway does not delete this object. If you are redeploying with a "
        "previously used location_name, choose a new, unique name instead.",
    ),
    (
        "Cannot create gateway: pending activation changes",
        "A prior configuration change in your Zscaler tenant has not yet been "
        "activated. Activate pending changes via the Zscaler Admin Portal, then "
        "retry the deployment.",
    ),
    (
        "Only one subnet can be specified for a Gateway Load Balancer endpoint",
        "This is expected AWS behavior and is already handled by this "
        "deployment — one endpoint is created per subnet automatically. If you "
        "see this error, confirm you are using the latest version of this "
        "deployment package.",
    ),
    (
        "Route already exists",
        "The destination CIDR you are routing to the GWLB Endpoint already has "
        "a route in that table (commonly an existing Internet Gateway or NAT "
        "Gateway default route). Use a dedicated subnet and route table for "
        "inspected traffic, separate from existing production routes.",
    ),
    (
        "Gateway stuck in an initializing state",
        "Provisioning typically completes within 3–10 minutes. If a timeout "
        "occurs, simply re-run the deployment command — it will detect the "
        "existing gateway and resume monitoring its health.",
    ),
]

for title, body in issues:
    p = doc.add_paragraph()
    tighten(p, before=2, after=2)
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = ORANGE
    run.font.size = Pt(11)
    add_body(doc, body)


# -----------------------------------------------------------------------
# 10. SECURITY NOTES
# -----------------------------------------------------------------------
add_heading_tight(doc, "10. Security Notes", 1)

add_bullet(doc, " Never commit your filled-in configuration file to version control. "
                "It contains your Zscaler client secret in plain text.",
           bold_prefix="Credential Handling: ")
add_bullet(doc, " Your Zscaler client secret is stored in Terraform state in plain "
                "text, by design, to enable a single-command deployment. Treat your "
                "Terraform state file as sensitive, and restrict access accordingly.",
           bold_prefix="Terraform State: ")
add_bullet(doc, " Grant only the minimum AWS IAM permissions listed in the "
                "Prerequisites section to the credentials used for deployment.",
           bold_prefix="Least Privilege: ")
add_bullet(doc, " Consider storing your Terraform state in a secure, access-controlled "
                "remote backend (e.g. an encrypted S3 bucket with restricted access) "
                "rather than locally, especially for production environments.",
           bold_prefix="Remote State (Recommended): ")

add_spacer(doc, 20)
closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
tighten(closing, before=0, after=0)
run = closing.add_run("End of Document")
run.font.italic = True
run.font.color.rgb = GRAY

# -----------------------------------------------------------------------
# Footer with page numbers
# -----------------------------------------------------------------------
def add_page_number_footer(document):
    for section in document.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tighten(p, before=0, after=0)
        run = p.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY

add_page_number_footer(doc)

output_path = "/Users/anasshahid/Downloads/ztgw-automation-aws-customer/docs/customer-onboarding-guide.docx"
doc.save(output_path)
print(f"Saved: {output_path}")
